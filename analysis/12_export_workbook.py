"""Collect result tables into one workbook."""

EXCEL_OUT = f'{OUT_DIR}/AMYGDALA_PD_MANUSCRIPT.xlsx'
FILE_LIST = [
    ('participant_flow.csv', 'Participant_Flow'),
    ('followup_duration_summary.csv', 'Followup_Duration'),
    ('cs4_whole_amygdala_PD_vs_HC.csv', 'CS4_Whole_Amygdala_PDvHC'),
    ('cs4_sectors_PD_vs_HC.csv', 'CS4_Sectors_PDvHC'),
    ('long5a_fixed_slopes_PD_vs_HC.csv', 'Long5a_Fixed_Slopes'),
    ('long5b_whole_amygdala_interaction.csv', 'Long5b_Whole_Interaction'),
    ('long5b_sectors_interaction.csv', 'Long5b_Sectors_Interaction'),
    ('long5b_nuclei_interaction_posthoc.csv', 'Long5b_Nuclei_PostHoc'),
    ('demographics_table.csv', 'Demographics'),
    ('demographics_within_subtype_significance.csv', 'Demo_Within_Subtype'),
    ('PD_vs_HC_clinical.csv', 'PD_vs_HC_Clinical'),
    ('VH11b_subtype_sectors.csv', 'VH11b_Subtype_Sectors'),
    ('VH11b_subtype_sectors_pairwise.csv', 'VH11b_Subtype_Pairwise'),
    ('VH11b_subtype_nuclei_posthoc.csv', 'VH11b_Subtype_Nuclei_PH'),
    ('VH11c_longitudinal_subtype_time.csv', 'VH11c_Long_Subtype_Time'),
    ('VH_9D_sectors.csv', 'VH9D_Sectors_CS'),
    ('VH_9D_nuclei.csv', 'VH9D_Nuclei_PostHoc'),
    ('VH_9H_demographics.csv', 'VH9H_Demographics'),
    ('VH_9G_survival.csv', 'VH9G_Survival'),
    ('CI_9D_sectors.csv', 'CI9D_Sectors_CS'),
    ('CI_9D_nuclei.csv', 'CI9D_Nuclei_PostHoc'),
    ('CI_9H_demographics.csv', 'CI9H_Demographics'),
    ('CI_9G_survival.csv', 'CI9G_Survival'),
    ('VH_9G_cohort_flow.csv', 'VH9G_Cohort_Flow'),
    ('CI_9G_cohort_flow.csv', 'CI9G_Cohort_Flow'),
    ('VHge2_9D_sectors.csv', 'SensVH2_9D_Sectors'),
    ('VHge2_9G_survival.csv', 'SensVH2_9G_Survival'),
    ('Step9J_random_slope_sensitivity.csv', 'RandomSlope_Sensitivity'),
]

used_names = set()
def safe_sheet_name(name, used):
    name = name[:31]; base = name; i = 1
    while name in used:
        suffix = f'_{i}'; name = base[:31-len(suffix)] + suffix; i += 1
    used.add(name); return name

toc_rows = []
with pd.ExcelWriter(EXCEL_OUT, engine='openpyxl') as writer:
    for fname, sheet_name in FILE_LIST:
        fpath = f'{OUT_DIR}/{fname}'
        if not os.path.exists(fpath):
            print(f"  Missing (skipped): {fname}"); continue
        try:
            df_sheet = safe_read_csv(fpath)
        except Exception as e:
            continue
        final_name = safe_sheet_name(sheet_name, used_names)
        df_sheet.to_excel(writer, sheet_name=final_name, index=False)
        toc_rows.append({'Sheet':final_name,'Source_File':fname,'N_rows':len(df_sheet)})
        print(f"  Wrote sheet: {final_name:<31} ({len(df_sheet)} rows)")
    pd.DataFrame(toc_rows).to_excel(writer, sheet_name='INDEX', index=False)

from openpyxl import load_workbook
wb = load_workbook(EXCEL_OUT)
if 'INDEX' in wb.sheetnames:
    wb.move_sheet('INDEX', offset=-(len(wb.sheetnames)-1))
wb.save(EXCEL_OUT)

print(f"\nMaster workbook saved: {EXCEL_OUT}")
